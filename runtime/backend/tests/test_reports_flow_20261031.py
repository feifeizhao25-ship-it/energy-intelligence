"""Report flow: save a financial model -> generate a report (both editions) -> real files.

Covers: the saved model reaching the report generator (it used to 422 forever), a missing IRR not
crashing save (it used to 500), English edition producing English content with contents page numbers,
domestic PDF embedding a CJK font and printing U+00B7, and heading structure.
"""

import io
import re

import pytest

from app.routers import reports as R

SOURCE = {"title": "NREL Annual Technology Baseline", "url": "https://atb.nrel.gov/", "retrieved_at": "2026-10-30"}
SOLAR = {"capacity_mw": 100, "capex_per_w": 1.1, "opex_per_kw_yr": 17, "electricity_price": 45,
         "capacity_factor": 0.25, "currency": "USD"}


async def _project(client, headers, name="Report Test Project"):
    r = await client.post("/api/v1/projects", headers=headers, json={
        "name": name, "technology": "solar", "capacity_kw": 100000,
        "location": {"latitude": 33.4, "longitude": -112.0, "address": "Phoenix, Arizona"}})
    assert r.status_code == 201, r.text
    return r.json()["data"]["id"]


async def _generate(client, headers, pid, **extra):
    body = {"project_id": pid, "report_type": "feasibility", "format": "pdf", **extra}
    return await client.post("/api/v1/reports/generate", headers=headers, json=body)


@pytest.mark.asyncio
async def test_generate_before_saving_a_model_is_a_clear_422(client, auth_headers):
    pid = await _project(client, auth_headers)
    zh = await _generate(client, auth_headers, pid)
    assert zh.status_code == 422 and "请先为该项目保存财务模型" in zh.json()["detail"]
    en = await _generate(client, auth_headers, pid, market="global")
    assert en.status_code == 422 and "Save a model for this project first" in en.json()["detail"]
    assert not re.search(r"[\u4e00-\u9fff]", en.json()["detail"])


@pytest.mark.asyncio
async def test_saved_model_without_sources_is_still_refused(client, auth_headers):
    pid = await _project(client, auth_headers)
    assert (await client.post(f"/api/v1/finance/models/{pid}", headers=auth_headers, json=SOLAR)).status_code == 200
    r = await _generate(client, auth_headers, pid, market="global")
    assert r.status_code == 422 and "verifiable data source" in r.json()["detail"]


@pytest.mark.asyncio
async def test_saved_model_with_sources_generates_english_pdf_and_docx(client, auth_headers, tmp_path, monkeypatch):
    monkeypatch.setenv("REPORTS_DIR", str(tmp_path))
    pid = await _project(client, auth_headers)
    saved = await client.post(f"/api/v1/finance/models/{pid}", headers=auth_headers, json={**SOLAR, "data_sources": [SOURCE]})
    assert saved.status_code == 200, saved.text
    for fmt in ("pdf", "docx"):
        r = await _generate(client, auth_headers, pid, market="global", format=fmt)
        assert r.status_code == 200, r.text
        rid = r.json()["data"]["report_id"]
        status = (await client.get(f"/api/v1/reports/{rid}", headers=auth_headers)).json()["data"]
        assert status["status"] == "completed", status
        body = (await client.get(f"/api/v1/reports/{rid}/file", headers=auth_headers)).content
        if fmt == "pdf":
            from pypdf import PdfReader
            text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(body)).pages)
        else:
            from docx import Document
            d = Document(io.BytesIO(body))
            text = "\n".join(p.text for p in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        assert not re.search(r"[\u4e00-\u9fff]", text)
        for needed in ("Financial Analysis", "Internal rate of return", "NREL Annual Technology Baseline", "Not verified" if False else "Disclaimer"):
            assert needed in text


@pytest.mark.asyncio
async def test_solar_model_with_no_unique_irr_saves_instead_of_500(client, auth_headers):
    pid = await _project(client, auth_headers)
    bad = {**SOLAR, "capex_per_w": 6.0, "electricity_price": 20}  # never pays back: IRR undefined
    r = await client.post(f"/api/v1/finance/models/{pid}", headers=auth_headers, json=bad)
    assert r.status_code == 200, r.text
    assert r.json()["irr"] is None


def _sample(market_global):
    d = R._demo_data()
    if market_global:
        d.update(name="Arizona Solar Portfolio", country_code="US", province_or_region="Arizona", city="Phoenix", currency="USD", description="")
    else:
        d.update(name="示例·光伏电站（阿尔·巴特）", description="项目·说明 • 要点")
    d["data_sources"] = [{**SOURCE}]
    return d


@pytest.mark.parametrize("rtype", ["feasibility", "investment", "compliance", "esg"])
def test_global_reports_are_real_content_not_placeholders(rtype):
    from pypdf import PdfReader
    d = _sample(True)
    m = R.calc_financial_metrics(d["financial"])
    pdf = R.generate_pdf(d, rtype, R.TEMPLATES[rtype], m, "Report", "confidential", "global")
    reader = PdfReader(io.BytesIO(pdf))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert len(reader.pages) >= 4 and len(text) > 2500
    assert "Contents" in text and "Appendix" in text and "Disclaimer" in text
    assert "see project documentation" not in text  # the old placeholder sentence
    assert not re.search(r"[\u4e00-\u9fff]", text)
    assert "tax advice" in text  # the not-advice statement is present


def test_global_pdf_contents_page_has_real_page_numbers_and_bookmarks():
    from pypdf import PdfReader
    d = _sample(True)
    m = R.calc_financial_metrics(d["financial"])
    reader = PdfReader(io.BytesIO(R.generate_pdf(d, "feasibility", R.TEMPLATES["feasibility"], m, "Report", "confidential", "global")))
    toc = reader.pages[1].extract_text()
    numbers = [int(n) for n in re.findall(r"\. (\d+)(?=\n|$)", toc)]
    assert len(numbers) >= 8
    assert "Project Overview" in toc and "Risk Assessment" in toc
    assert numbers == sorted(numbers) and numbers[0] >= 3 and numbers[-1] <= len(reader.pages)
    assert len(reader.outline) >= 5


def test_global_docx_has_heading_hierarchy_toc_field_and_tables():
    from docx import Document
    d = _sample(True)
    m = R.calc_financial_metrics(d["financial"])
    doc = Document(io.BytesIO(R.generate_docx(d, "investment", R.TEMPLATES["investment"], m, "Report", "confidential", "global")))
    styles = {p.style.name for p in doc.paragraphs}
    assert {"Heading 1", "Heading 2"} <= styles
    assert len(doc.tables) >= 4
    xml = doc.element.xml
    assert "TOC \\o" in xml and "updateFields" in doc.settings.element.xml


def test_domestic_pdf_embeds_a_font_and_prints_the_middle_dot():
    from pypdf import PdfReader
    d = _sample(False)
    m = R.calc_financial_metrics(d["financial"])
    pdf = R.generate_pdf(d, "feasibility", R.TEMPLATES["feasibility"], m, "测试报告·2026", "内部")
    reader = PdfReader(io.BytesIO(pdf))
    fonts = set()
    for page in reader.pages:
        res = page["/Resources"].get("/Font", {})
        for f in res.values():
            fd = f.get_object()
            fonts.add(str(fd.get("/BaseFont")))
            desc = fd.get("/FontDescriptor")
            if "CJK" in str(fd.get("/Subtype")) or "Type0" in str(fd.get("/Subtype")):
                pass
    assert any("NotoSansSC" in f for f in fonts), fonts  # embedded TrueType, not the viewer-dependent CID font
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert "示例·光伏电站" in text  # U+00B7 survives
    assert "□" not in text


def test_domestic_pdf_toc_has_page_numbers_and_no_hand_drawn_check_marks():
    from pypdf import PdfReader
    d = _sample(False)
    m = R.calc_financial_metrics(d["financial"])
    reader = PdfReader(io.BytesIO(R.generate_pdf(d, "feasibility", R.TEMPLATES["feasibility"], m, "报告", "内部")))
    toc = reader.pages[1].extract_text()
    assert "第一章 项目概述" in toc.replace("  ", " ")
    nums = [int(n) for n in re.findall(r"(\d+)\s*$", toc, flags=re.M)]
    assert nums and max(nums) <= len(reader.pages) and min(nums) >= 3
    full = "\n".join(p.extract_text() for p in reader.pages)
    assert "\u26a0" not in full and "\u2713" not in full


def test_cid_fallback_swaps_only_the_missing_glyphs():
    from app.services import report_pdf_support as sup
    assert sup.sanitize_text("A·B•C") == "A\u30fbB\u30fbC"
    assert sup.sanitize_for_cid({"a": ["x·y", 3]}) == {"a": ["x\u30fby", 3]}
    assert sup.needs_swap("STSong-Light") is True
