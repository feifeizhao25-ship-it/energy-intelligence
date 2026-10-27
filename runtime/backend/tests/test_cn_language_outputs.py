"""中文输出扫描：覆盖率清单，以及导出样例不得夹带英文句子。"""
from __future__ import annotations

import json
import subprocess
from io import BytesIO
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[3]
SAMPLES = ROOT / "docs" / "report-samples" / "cn"
WEB = ROOT / "runtime" / "web"
BACKEND = ROOT / "runtime" / "backend"


def test_output_language_coverage_lists_scanned_paths():
    completed = subprocess.run(
        ["node", "scripts/check-output-language.mjs"],
        cwd=WEB,
        check=True,
        capture_output=True,
        text=True,
    )
    report = json.loads(completed.stdout)
    scanned = "\n".join(report["scanned"])
    assert "src/lib/reports" in scanned or any("reports/" in item for item in report["scanned"])
    assert "src/lib/assistant" in scanned or any("assistant/" in item for item in report["scanned"])
    assert any(item.endswith(".md") for item in report["scanned"])
    assert report["sampleViolations"] == []
    detail_files = [
        str(path.relative_to(BACKEND))
        for path in (BACKEND / "app").rglob("*.py")
        if "detail=" in path.read_text(encoding="utf-8") or "message=" in path.read_text(encoding="utf-8")
    ]
    assert detail_files, "应当扫到带 detail 或 message 的后端文件"
    assert any("sms" in name or "auth" in name for name in detail_files)


def test_three_export_samples_have_no_english_prose():
    names = sorted(SAMPLES.glob("*.md"))[:3]
    assert len(names) == 3
    banned = ("please ", " the ", "error", "welcome")
    for path in names:
        text = path.read_text(encoding="utf-8")
        document = Document()
        document.add_paragraph(text)
        buffer = BytesIO()
        document.save(buffer)
        loaded = Document(BytesIO(buffer.getvalue()))
        plain = "\n".join(paragraph.text for paragraph in loaded.paragraphs).lower()
        assert "示例数据" in plain
        for word in banned:
            assert word not in plain
