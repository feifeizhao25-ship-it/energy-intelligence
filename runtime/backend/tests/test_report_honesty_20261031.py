"""Domestic compliance reports must not claim "符合"; English evidence text must be English."""
import io
import json
import re

import pytest

from app.routers import reports as R


def _data(sources=()):
    d = R._demo_data()
    d.update(province_or_region="新疆", name="示例光伏电站")
    d["data_sources"] = list(sources)
    return d


def _docx_text(d):
    from docx import Document
    m = R.calc_financial_metrics(d["financial"])
    doc = Document(io.BytesIO(R.generate_docx(d, "compliance", R.TEMPLATES["compliance"], m, "合规报告", "confidential", "cn")))
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_domestic_compliance_report_makes_no_compliance_claim_without_evidence():
    text = _docx_text(_data())
    assert "符合" not in text.replace("符合性须人工确认", "")
    assert "✓" not in text and "√" not in text
    assert "未登记来源，待核对" in text


def test_registered_source_changes_the_row_and_still_needs_a_person():
    src = {"title": "《电力法》全文", "url": "https://flk.npc.gov.cn/x", "publisher": "全国人大"}
    rows = {r[0]: r[2] for r in R.compliance_rows(_data([src]))}
    assert rows["《电力法》"] == "已登记来源，待人工确认"
    assert rows["《可再生能源法》"] == "未登记来源，待核对"
    assert rows["地方新能源规划"] == "未登记来源，待核对"
    assert all("符合" not in status for status in rows.values())


def test_missing_province_is_called_out_not_passed():
    d = _data()
    d["province_or_region"] = ""
    row = [r for r in R.compliance_rows(d) if r[0] == "地方新能源规划"][0]
    assert "无法核对" in row[2]


def test_domestic_compliance_pdf_has_no_check_marks():
    pytest.importorskip("reportlab")
    from pypdf import PdfReader
    d = _data()
    m = R.calc_financial_metrics(d["financial"])
    reader = PdfReader(io.BytesIO(R.generate_pdf(d, "compliance", R.TEMPLATES["compliance"], m, "合规报告", "confidential", "cn")))
    text = "\n".join(p.extract_text() for p in reader.pages)
    assert "待人工确认" in text or "待核对" in text
    assert "√ 符合" not in text and "✓ 符合" not in text


def test_english_evidence_envelope_has_no_chinese():
    d = R._demo_data()
    d["data_sources"] = [{"title": "Grid tariff notice", "url": "https://example.gov/a", "publisher": "State regulator"}]
    env = R._build_report_evidence("Feasibility - 2026-10-31", d, True)
    blob = json.dumps(env, ensure_ascii=False)
    assert not re.search(r"[\u3400-\u9fff]", blob), re.findall(r"[^\"]*[\u3400-\u9fff][^\"]*", blob)[:5]
    assert env["answer"].startswith("Report job accepted")


def test_chinese_evidence_envelope_stays_chinese():
    env = R._build_report_evidence("可研报告", R._demo_data(), False)
    assert env["answer"].startswith("报告生成任务已受理")
