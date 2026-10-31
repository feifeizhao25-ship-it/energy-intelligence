import sys, json, re, os, subprocess, io
sys.path.insert(0, '/workspace/ui_full_test/run')
from enapi import call
exec(open('/workspace/ui_full_test/run/en_battery.py').read().split('def run(pid):')[0].split('P={',1)[0].replace('import sys, json, time\nfrom enapi import *',''))
src = open('/workspace/ui_full_test/run/en_battery.py').read()
P = eval('{' + src.split('P={',1)[1].split('\ndef run(pid):')[0])
PORT = {'cn': 4291, 'global': 4292}
OUT = '/workspace/ui_full_test/retest/after'
CJK = re.compile(r'[\u3400-\u9fff]')
res = {}
def text_of(path, fmt):
    if fmt == 'pdf':
        return subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True).stdout
    from docx import Document
    d = Document(path)
    return '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
for pid, c in P.items():
    port = PORT[c['market']]
    r = {'persona': c['name'], 'market': c['market']}
    pw = 'Test@12345Aa'
    call(port, 'POST', '/api/v1/auth/register', dict(name=c['name'], email=c['email'], password=pw, company=c['company'] or None, role=c['role'], country=c['country'], market=c['market'], data_transfer_consent=True))
    st, j, _ = call(port, 'POST', '/api/v1/auth/login', {'email': c['email'], 'password': pw})
    tok = (j.get('data', j) or {}).get('access_token') if isinstance(j, dict) else None
    if not tok: r['login'] = 'FAILED'; res[pid] = r; continue
    lat, lng, addr = c['loc']
    st, pj, _ = call(port, 'POST', '/api/v1/projects', {'name': f"{c['name']} project", 'location': {'latitude': lat, 'longitude': lng, 'address': addr}, 'capacity_kw': c['cap'], 'technology': c['tech']}, tok)
    pd = pj.get('data', pj); pidd = pd.get('id') or pd.get('project_id'); r['project'] = st
    model = c.get('solar') or c.get('wind')
    if not model:
        r['note'] = 'storage persona: /finance/models accepts solar/wind only, so no report can be generated (unchanged limitation)'
        res[pid] = r; continue
    st, m, _ = call(port, 'POST', f'/api/v1/finance/models/{pidd}', model, tok); r['save_model'] = st
    st, b, _ = call(port, 'POST', '/api/v1/reports/generate', {'project_id': pidd, 'report_type': 'feasibility', 'format': 'pdf', 'market': c['market']}, tok)
    r['report_without_sources'] = {'status': st, 'detail': (b.get('detail') if isinstance(b, dict) else str(b))[:160]}
    hdr = {}
    src_list = [{'title': 'Regional grid tariff notice' if c['market'] == 'global' else '地方电价通知', 'url': 'https://example.gov/tariff', 'publisher': 'Regulator' if c['market'] == 'global' else '地方发改委', 'published_at': '2026-09-01'}]
    st, b, _ = call(port, 'PUT', f'/api/v1/finance/data-sources/{pidd}', {'sources': src_list}, tok); r['put_sources'] = {'status': st, 'count': b.get('count') if isinstance(b, dict) else None}
    for rtype in ('feasibility', 'compliance'):
        for fmt in ('pdf', 'docx'):
            st, b, _ = call(port, 'POST', '/api/v1/reports/generate', {'project_id': pidd, 'report_type': rtype, 'format': fmt, 'market': c['market']}, tok)
            key = f'{rtype}_{fmt}'
            if st != 200: r[key] = {'status': st, 'detail': str(b)[:200]}; continue
            rid = b['data']['report_id']
            st2, body, h, _ = call(port, 'GET', f'/api/v1/reports/{rid}/file', None, tok, raw=True)
            path = f'{OUT}/en_{pid}_{rtype}.{fmt}'; open(path, 'wb').write(body)
            t = text_of(path, fmt)
            r[key] = {'status': st2, 'chars': len(t), 'has_cjk': bool(CJK.search(t)), 'has_fuhe_claim': ('符合' in t.replace('符合性须人工确认', '')),
                      'has_check_mark': ('✓' in t or '√' in t)}
            if fmt == 'pdf':
                subprocess.run(['pdftoppm', '-png', '-r', '60', '-f', '1', '-l', '6', path, f'{OUT}/en_{pid}_{rtype}_p'])
    res[pid] = r
json.dump(res, open(f'{OUT}/en_retest_summary.json', 'w'), ensure_ascii=False, indent=1)
for k, v in res.items(): print(k, json.dumps(v, ensure_ascii=False)[:700])
